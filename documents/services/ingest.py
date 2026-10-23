"""Bounded, local text acquisition for PDF and DOCX documents.

No document contents are sent over the network. PDF scans are rendered and OCR'd
one page at a time, so a long document does not create parallel OCR processes.
DOCX files have no dependable page numbers without a layout engine: their
sources are labelled as paragraphs and table rows instead.
"""

from __future__ import annotations

import os
import re
import shutil
import subprocess
import tempfile
from pathlib import Path
from zipfile import ZipFile

from docx import Document
from docx.table import Table
from pypdf import PdfReader


class IngestionError(ValueError):
    """A document could not be read within the supported input constraints."""


MAX_FILE_BYTES = 50 * 1024 * 1024
MAX_DOCX_EXPANDED_BYTES = 100 * 1024 * 1024
OCR_DPI = 180
OCR_MAX_SIDE_PIXELS = 2200
_WORD_OMITTED_TAGS = re.compile(
    rb"<(?:[A-Za-z_][A-Za-z0-9_.-]*:)?"
    rb"(txbxContent|ins|del|moveFrom|moveTo|object|OLEObject)(?=[\s/>])"
)


def _clean(text: str) -> str:
    return text.replace("\x00", "").replace("\r\n", "\n").strip()


def _ocr_page(path: Path, page_number: int, language: str) -> str:
    """Return OCR text without exposing subprocess diagnostics or input text."""
    renderer = shutil.which("pdftoppm")
    tesseract = shutil.which("tesseract")
    if not renderer or not tesseract:
        raise IngestionError(
            "Il PDF contiene pagine scansionate: installa Poppler (pdftoppm) "
            "e Tesseract con i dati della lingua richiesta."
        )
    env = os.environ.copy()
    env["OMP_THREAD_LIMIT"] = "1"
    env["OMP_NUM_THREADS"] = "1"
    try:
        with tempfile.TemporaryDirectory(prefix="document-ai-ocr-") as temporary:
            prefix = Path(temporary) / "page"
            subprocess.run(
                [
                    renderer, "-f", str(page_number), "-l", str(page_number),
                    "-singlefile", "-r", str(OCR_DPI), "-scale-to",
                    str(OCR_MAX_SIDE_PIXELS), "-png", str(path), str(prefix),
                ],
                check=True, stdout=subprocess.DEVNULL, stderr=subprocess.PIPE,
                timeout=90, env=env,
            )
            image = prefix.with_suffix(".png")
            if not image.is_file():
                raise IngestionError("La conversione della pagina PDF non ha prodotto un'immagine.")
            completed = subprocess.run(
                [tesseract, str(image), "stdout", "-l", language, "--psm", "3"],
                check=True, stdout=subprocess.PIPE, stderr=subprocess.PIPE,
                timeout=120, env=env, text=True, encoding="utf-8", errors="replace",
            )
            return _clean(completed.stdout)
    except subprocess.TimeoutExpired as exc:
        raise IngestionError(
            f"Tempo massimo OCR superato alla pagina {page_number}. "
            "Prova un PDF più breve o una scansione meno pesante."
        ) from exc
    except subprocess.CalledProcessError as exc:
        raise IngestionError(
            f"OCR non riuscito alla pagina {page_number}. Verifica la leggibilità "
            f"del PDF e la disponibilità delle lingue Tesseract: {language}."
        ) from exc
    except OSError as exc:
        raise IngestionError("Impossibile eseguire gli strumenti OCR locali.") from exc


def _has_raster_image(resources, depth: int = 0) -> bool:
    """Check image XObjects, including the common nested-Form case."""
    if not resources or depth > 4:
        return False
    try:
        resources = resources.get_object()
        xobjects = resources.get("/XObject", {})
        xobjects = xobjects.get_object() if hasattr(xobjects, "get_object") else xobjects
        for reference in xobjects.values():
            obj = reference.get_object()
            if obj.get("/Subtype") == "/Image":
                return True
            if obj.get("/Subtype") == "/Form" and _has_raster_image(obj.get("/Resources"), depth + 1):
                return True
    except (AttributeError, TypeError, ValueError, KeyError):
        # Empty or malformed resource metadata must not discard readable text.
        return False
    return False


def _read_pdf(path: Path, language: str, max_pages: int, add, warnings: list[str]) -> int:
    try:
        reader = PdfReader(path, strict=False)
        if reader.is_encrypted:
            raise IngestionError("Il PDF è protetto da password. Carica una copia non protetta.")
        page_count = len(reader.pages)
        if page_count > max_pages:
            raise IngestionError(f"Il PDF supera il limite di {max_pages} pagine.")
        if not page_count:
            raise IngestionError("Il PDF non contiene pagine.")
        for number, page in enumerate(reader.pages, start=1):
            try:
                text = _clean(page.extract_text(extraction_mode="layout") or "")
            except Exception:
                try:
                    text = _clean(page.extract_text() or "")
                    warnings.append(f"Pagina {number}: usata l'estrazione testuale semplice.")
                except Exception:
                    text = ""
            method = "pdf_text"
            compact = re.sub(r"\s", "", text)
            # Image pages can contain just a page number or a scanner watermark.
            if not compact or (len(compact) < 40 and _has_raster_image(page.get("/Resources"))):
                text = _ocr_page(path, number, language)
                method = "ocr"
            if not text:
                warnings.append(f"Pagina {number}: nessun testo leggibile rilevato.")
            add({"id": f"p{number}", "label": f"Pagina {number}", "page": number,
                 "text": text, "method": method})
        return page_count
    except IngestionError:
        raise
    except Exception as exc:
        raise IngestionError("Impossibile leggere il PDF: il file potrebbe essere danneggiato.") from exc


def _read_docx(path: Path, add, warnings: list[str]) -> None:
    try:
        with ZipFile(path) as archive:
            entries = archive.infolist()
            if sum(item.file_size for item in entries) > MAX_DOCX_EXPANDED_BYTES:
                raise IngestionError("Il documento Word supera il limite di contenuto decompresso.")
            contains_images = any(item.filename.startswith("word/media/") for item in entries)
            filenames = {item.filename for item in entries}
            # python-docx's body iterator does not expose these containers.
            # Check the bounded ZIP parts explicitly rather than silently
            # treating an incomplete acquisition as full Word coverage.
            omitted_tags = set()
            for item in entries:
                if item.filename.startswith("word/") and item.filename.endswith(".xml"):
                    omitted_tags.update(_WORD_OMITTED_TAGS.findall(archive.read(item)))
            if "word/footnotes.xml" in filenames:
                warnings.append(
                    "Il Word contiene note a piè di pagina: il loro contenuto non viene acquisito. "
                    "Per includerle, esporta il documento in PDF e carica il PDF."
                )
            if "word/endnotes.xml" in filenames:
                warnings.append(
                    "Il Word contiene note di chiusura: il loro contenuto non viene acquisito. "
                    "Per includerle, esporta il documento in PDF e carica il PDF."
                )
            if b"txbxContent" in omitted_tags:
                warnings.append(
                    "Il Word contiene caselle di testo: il testo al loro interno non viene acquisito. "
                    "Per includerlo, esporta il documento in PDF e carica il PDF."
                )
            if omitted_tags.intersection({b"ins", b"del", b"moveFrom", b"moveTo"}):
                warnings.append(
                    "Il Word contiene revisioni (modifiche tracciate): il testo inserito, eliminato "
                    "o spostato potrebbe non essere acquisito. Accetta o rifiuta le revisioni "
                    "in una copia del documento prima di elaborarlo."
                )
            if omitted_tags.intersection({b"object", b"OLEObject"}) or any(
                filename.startswith("word/embeddings/") for filename in filenames
            ):
                warnings.append(
                    "Il Word contiene oggetti incorporati: il loro contenuto non viene acquisito. "
                    "Carica separatamente i documenti incorporati o una copia PDF che li includa."
                )
        document = Document(path)
        if contains_images:
            warnings.append(
                "Il Word contiene immagini: l'OCR delle immagini incorporate non è supportato. "
                "Per acquisirle, esporta il documento in PDF e carica il PDF."
            )
        warnings.append(
            "Le fonti Word indicano paragrafi e righe di tabella: "
            "la numerazione delle pagine non è disponibile nel formato DOCX."
        )
        serial = 0

        def append(text: str, label: str):
            nonlocal serial
            text = _clean(text)
            if text:
                serial += 1
                add({"id": f"b{serial}", "label": label, "page": None,
                     "text": text, "method": "docx"})

        # Word stores headers and footers separately from body paragraphs.
        # Linked section headers share the same element; include each only once.
        seen_parts = set()
        for section_number, section in enumerate(document.sections, start=1):
            for label, part in (
                ("Intestazione", section.header),
                ("Intestazione prima pagina", section.first_page_header),
                ("Intestazione pagine pari", section.even_page_header),
            ):
                marker = part._element
                if marker in seen_parts:
                    continue
                seen_parts.add(marker)
                for paragraph in part.paragraphs:
                    append(paragraph.text, f"{label}, sezione {section_number}")
                for table in part.tables:
                    for row_number, row in enumerate(table.rows, start=1):
                        append(" | ".join(cell.text for cell in row.cells),
                               f"{label}, sezione {section_number}, riga {row_number}")

        paragraph_number = table_number = 0
        for block in document.iter_inner_content():
            if isinstance(block, Table):
                table_number += 1
                for row_number, row in enumerate(block.rows, start=1):
                    cells = []
                    previous_cell = None
                    for cell in row.cells:
                        if cell._tc is not previous_cell:
                            cells.append(cell.text.strip())
                        previous_cell = cell._tc
                    append(" | ".join(cells), f"Tabella {table_number}, riga {row_number}")
            else:
                paragraph_number += 1
                append(block.text, f"Paragrafo {paragraph_number}")

        for section_number, section in enumerate(document.sections, start=1):
            for label, part in (
                ("Piè di pagina", section.footer),
                ("Piè di prima pagina", section.first_page_footer),
                ("Piè di pagine pari", section.even_page_footer),
            ):
                marker = part._element
                if marker in seen_parts:
                    continue
                seen_parts.add(marker)
                for paragraph in part.paragraphs:
                    append(paragraph.text, f"{label}, sezione {section_number}")
                for table in part.tables:
                    for row_number, row in enumerate(table.rows, start=1):
                        append(" | ".join(cell.text for cell in row.cells),
                               f"{label}, sezione {section_number}, riga {row_number}")
    except IngestionError:
        raise
    except Exception as exc:
        raise IngestionError("Impossibile leggere il Word: carica un file DOCX valido, non protetto.") from exc


def extract_document(
    path: Path, *, ocr_language: str = "ita+eng", max_pages: int = 50,
    max_chars: int = 250000,
) -> dict:
    """Extract a PDF/DOCX into source-addressable segments, preserving originals.

    Limits reject the job instead of silently presenting a truncated document as
    complete. DOCX pagination cannot be inferred: ``max_pages`` applies to PDF;
    Word inputs are bounded by file, expanded ZIP and extracted text sizes.
    """
    path = Path(path)
    if max_pages < 1 or max_chars < 1:
        raise IngestionError("I limiti di pagine e caratteri devono essere positivi.")
    if not re.fullmatch(r"[A-Za-z0-9_]+(?:\+[A-Za-z0-9_]+)*", ocr_language):
        raise IngestionError("Configurazione della lingua OCR non valida.")
    if path.suffix.lower() not in {".pdf", ".docx"}:
        raise IngestionError("Formato non supportato. Carica un PDF oppure un Word DOCX.")
    try:
        if not path.is_file():
            raise IngestionError("Il documento non è disponibile.")
        if path.stat().st_size > MAX_FILE_BYTES:
            raise IngestionError("Il file supera il limite di 50 MB.")
    except OSError as exc:
        raise IngestionError("Impossibile accedere al documento.") from exc

    segments: list[dict] = []
    warnings: list[str] = []
    count = 0

    def add(segment: dict):
        nonlocal count
        count += len(segment["text"])
        if count > max_chars:
            raise IngestionError(f"Il testo supera il limite di {max_chars:,} caratteri.")
        segments.append(segment)

    page_count = None
    if path.suffix.lower() == ".pdf":
        page_count = _read_pdf(path, ocr_language, max_pages, add, warnings)
    else:
        _read_docx(path, add, warnings)
    if not any(segment["text"].strip() for segment in segments):
        raise IngestionError("Nessun testo leggibile trovato nel documento.")
    return {"segments": segments, "warnings": list(dict.fromkeys(warnings)),
            "page_count": page_count, "text_char_count": count}
