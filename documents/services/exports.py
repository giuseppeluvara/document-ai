import json
from io import BytesIO
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

def effective_fields(document):
    definitions = {item['key']: item for item in document.template_snapshot.get('fields', [])}
    result = []
    for field in document.result.get('fields', []):
        item = dict(field)
        key = item['key']
        item['type'] = definitions.get(key, {}).get('type', 'text')
        item['original_value'] = field.get('value')
        item['automatic_status'] = field.get('status')
        item['manually_reviewed'] = key in document.overrides
        if item['manually_reviewed']:
            item['value'] = document.overrides[key]['value']
            item['reviewed_at'] = document.overrides[key]['reviewed_at']
            item['status'] = 'not_found' if item['value'] is None else 'found'
        result.append(item)
    return result

def export_payload(document):
    return {
        'schema_version': '1.0',
        'document': {'id': str(document.id), 'filename': document.original_name, 'sha256': document.sha256, 'created_at': document.created_at.isoformat()},
        'template': document.template_snapshot,
        'engine': document.result.get('engine'),
        'fields': effective_fields(document),
        'sources': document.extracted.get('segments', []),
        'warnings': document.extracted.get('warnings', []) + document.result.get('warnings', []),
        'metrics': document.metrics,
        'note': 'Le citazioni descrivono il risultato automatico. Le modifiche manuali sono conservate separatamente e non sono validate dal modello.',
    }

def _display(value):
    if value is None:
        return ''
    if isinstance(value, (dict, list)):
        return json.dumps(value, ensure_ascii=False)
    return str(value)

def build_xlsx(document):
    wb = Workbook()
    ws = wb.active
    ws.title = 'Campi'
    rows = [['Chiave', 'Campo', 'Tipo', 'Valore', 'Valore automatico', 'Revisione manuale', 'Fonti', 'Avvisi']]
    for field in effective_fields(document):
        rows.append([field['key'], field['label'], field['type'], _display(field.get('value')),
                     _display(field.get('original_value')), 'Sì' if field['manually_reviewed'] else 'No',
                     '\n'.join(f'{e["segment_id"]}: {e["quote"]}' for e in field.get('evidence', [])),
                     '\n'.join(field.get('warnings', []))])
    def append_full(sheet, row):
        """Excel caps cells at 32767 chars: write explicit continuation rows."""
        row = [''.join(f'\\u{ord(c):04x}' if (ord(c) < 32 and c not in '\t\n\r') or ord(c) in {65534, 65535} else c for c in v) if isinstance(v, str) else v for v in row]
        pieces = [[v[i:i+30000] for i in range(0, len(v), 30000)] if isinstance(v, str) and v else [v] for v in row]
        for part in range(max(len(value) for value in pieces)):
            sheet.append([value[part] if part < len(value) else (row[0] if col == 0 else '') for col, value in enumerate(pieces)])
    for row in rows:
        append_full(ws, row)
    sources = wb.create_sheet('Fonti')
    sources.append(['ID', 'Riferimento', 'Pagina', 'Metodo', 'Testo'])
    for segment in document.extracted.get('segments', []):
        # Excel cells have a 32767-character limit. Ingestion usually produces
        # shorter segments; split long sources explicitly rather than truncate.
        text = segment['text']
        for offset in range(0, max(len(text), 1), 30000):
            append_full(sources, [segment['id'], segment['label'], segment.get('page'), segment['method'], text[offset:offset + 30000]])
    meta = wb.create_sheet('Metadati')
    for row in [['Documento', document.original_name], ['SHA256', document.sha256], ['Template', document.template_snapshot.get('name', '')], ['Motore', document.result.get('engine', '')], ['Elaborazione (s)', document.metrics.get('elapsed_seconds', '')], ['Avvisi', '\n'.join(export_payload(document)['warnings'])], ['Nota', export_payload(document)['note']], ['Formato Excel', 'I testi oltre 30.000 caratteri continuano su righe successive. I caratteri di controllo non ammessi in XML sono rappresentati come \\uXXXX.']]:
        append_full(meta, row)
    for sheet in wb:
        sheet.freeze_panes = 'A2'
        sheet.auto_filter.ref = sheet.dimensions
        for row in sheet:
            for cell in row:
                if isinstance(cell.value, str):
                    # Keep values literal: prevent spreadsheet formula injection.
                    cell.data_type = 's'
                cell.alignment = Alignment(vertical='top', wrap_text=True)
        for cell in sheet[1]:
            cell.fill = PatternFill('solid', fgColor='183B56')
            cell.font = Font(color='FFFFFF', bold=True)
        for column in sheet.columns:
            sheet.column_dimensions[column[0].column_letter].width = 28 if column[0].column < 4 else 55
    output = BytesIO()
    wb.save(output)
    return output.getvalue()
