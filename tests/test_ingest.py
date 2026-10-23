"""Document reader tests use synthetic content, never customer documents."""

import io
import shutil
import subprocess
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch
from zipfile import ZipFile

from docx import Document
from PIL import Image, ImageDraw, ImageFont
from pypdf import PdfReader, PdfWriter

from documents.services.ingest import IngestionError, _ocr_page, extract_document

try:
    from reportlab.pdfgen import canvas
    from reportlab.lib.utils import ImageReader
except ImportError:
    canvas = None


class IngestionTests(unittest.TestCase):
    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory()
        self.root = Path(self.temporary.name)

    def tearDown(self):
        self.temporary.cleanup()

    def make_docx(self):
        document = Document()
        document.add_paragraph("Il contratto decorre dal 1 gennaio 2027.")
        table = document.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "Voce"
        table.cell(0, 1).text = "Importo"
        table.cell(1, 0).text = "Canone annuo"
        table.cell(1, 1).text = "EUR 12.000"
        document.add_paragraph("Durata: dodici mesi.")
        path = self.root / "contratto.docx"
        document.save(path)
        return path

    def test_docx_preserves_body_and_table_order_without_inventing_pages(self):
        result = extract_document(self.make_docx())
        self.assertIsNone(result["page_count"])
        self.assertEqual(
            [item["text"] for item in result["segments"]],
            ["Il contratto decorre dal 1 gennaio 2027.", "Voce | Importo",
             "Canone annuo | EUR 12.000", "Durata: dodici mesi."],
        )
        self.assertTrue(all(item["page"] is None for item in result["segments"]))
        self.assertEqual(result["segments"][2]["label"], "Tabella 1, riga 2")
        self.assertEqual(result["text_char_count"], sum(len(s["text"]) for s in result["segments"]))

    def test_docx_header_is_included_and_images_are_disclosed(self):
        path = self.make_docx()
        document = Document(path)
        document.sections[0].header.paragraphs[0].text = "Società di esempio S.r.l."
        image = io.BytesIO()
        Image.new("RGB", (20, 20), "white").save(image, format="PNG")
        image.seek(0)
        document.add_picture(image)
        document.save(path)
        result = extract_document(path)
        self.assertIn("Società di esempio", result["segments"][0]["text"])
        self.assertTrue(any("immagini" in warning for warning in result["warnings"]))

    def add_docx_parts(self, path, *, body_extra=b"", parts=None):
        with ZipFile(path) as archive:
            members = {name: archive.read(name) for name in archive.namelist()}
        members["word/document.xml"] = members["word/document.xml"].replace(
            b"</w:body>", body_extra + b"</w:body>",
        )
        members.update(parts or {})
        with ZipFile(path, "w") as archive:
            for name, data in members.items():
                archive.writestr(name, data)

    def test_docx_footnotes_and_textboxes_are_explicitly_reported_as_omitted(self):
        path = self.make_docx()
        self.add_docx_parts(
            path,
            body_extra=(
                b"<w:p><w:r><w:txbxContent><w:p><w:r><w:t>TESTO_NASCOSTO</w:t>"
                b"</w:r></w:p></w:txbxContent></w:r></w:p>"
            ),
            parts={
                "word/footnotes.xml": (
                    b'<w:footnotes xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
                    b'<w:footnote w:id="1"><w:p><w:r><w:t>NOTA_NON_ACQUISITA</w:t>'
                    b"</w:r></w:p></w:footnote></w:footnotes>"
                ),
            },
        )
        result = extract_document(path)
        warnings = " ".join(result["warnings"])
        self.assertIn("note a piè di pagina", warnings)
        self.assertIn("caselle di testo", warnings)
        acquired = " ".join(segment["text"] for segment in result["segments"])
        self.assertNotIn("TESTO_NASCOSTO", acquired)
        self.assertNotIn("NOTA_NON_ACQUISITA", acquired)

    def test_docx_endnotes_revisions_and_embedded_objects_have_coverage_warnings(self):
        path = self.make_docx()
        self.add_docx_parts(
            path,
            body_extra=(
                b'<w:ins w:id="1"><w:p><w:r><w:t>REVISIONE</w:t></w:r></w:p></w:ins>'
                b"<w:p><w:r><w:object/></w:r></w:p>"
            ),
            parts={
                "word/endnotes.xml": (
                    b'<w:endnotes xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"/>'
                ),
                "word/embeddings/embedded.bin": b"SYNTHETIC_EMBEDDED_OBJECT",
            },
        )
        result = extract_document(path)
        warnings = " ".join(result["warnings"])
        self.assertIn("note di chiusura", warnings)
        self.assertIn("modifiche tracciate", warnings)
        self.assertIn("oggetti incorporati", warnings)
        self.assertTrue(all(segment["method"] == "docx" for segment in result["segments"]))

    def test_character_limit_rejects_instead_of_silently_truncating(self):
        with self.assertRaisesRegex(IngestionError, "caratteri"):
            extract_document(self.make_docx(), max_chars=20)

    def test_unsupported_and_corrupt_inputs_fail_cleanly(self):
        unsupported = self.root / "contratto.doc"
        unsupported.write_text("content")
        with self.assertRaisesRegex(IngestionError, "Formato non supportato"):
            extract_document(unsupported)
        invalid = self.root / "contratto.docx"
        invalid.write_bytes(b"not a zip")
        with self.assertRaisesRegex(IngestionError, "DOCX valido"):
            extract_document(invalid)

    def test_docx_zip_expansion_limit(self):
        with patch("documents.services.ingest.MAX_DOCX_EXPANDED_BYTES", 10):
            with self.assertRaisesRegex(IngestionError, "decompresso"):
                extract_document(self.make_docx())

    def test_pdf_page_limit_and_password_are_rejected(self):
        path = self.root / "pages.pdf"
        writer = PdfWriter()
        writer.add_blank_page(width=200, height=200)
        writer.add_blank_page(width=200, height=200)
        writer.write(path)
        with self.assertRaisesRegex(IngestionError, "pagine"):
            extract_document(path, max_pages=1)
        writer.encrypt("test-password")
        writer.write(path)
        with self.assertRaisesRegex(IngestionError, "password"):
            extract_document(path)

    @unittest.skipUnless(canvas is not None, "reportlab is an optional test dependency")
    def test_text_pdf_does_not_need_ocr(self):
        path = self.root / "text.pdf"
        pdf = canvas.Canvas(str(path))
        pdf.drawString(30, 700, "Contratto di esempio: durata dodici mesi e canone annuo EUR 12000.")
        pdf.save()
        with patch("documents.services.ingest._ocr_page") as ocr:
            result = extract_document(path)
        ocr.assert_not_called()
        self.assertEqual(result["page_count"], 1)
        self.assertIn("canone annuo", result["segments"][0]["text"])
        self.assertEqual(result["segments"][0]["method"], "pdf_text")

    @unittest.skipUnless(canvas is not None, "reportlab is an optional test dependency")
    def test_mixed_pdf_ocr_only_scanned_page_and_preserves_source_pages(self):
        path = self.root / "mixed.pdf"
        pdf = canvas.Canvas(str(path))
        pdf.drawString(30, 700, "Prima pagina testuale del contratto con tutti i dati di esempio.")
        pdf.showPage()
        pdf.drawImage(ImageReader(Image.new("RGB", (100, 100), "white")), 10, 10, 200, 200)
        pdf.save()
        with patch("documents.services.ingest._ocr_page", return_value="Importo EUR 500,00") as ocr:
            result = extract_document(path)
        ocr.assert_called_once_with(path, 2, "ita+eng")
        self.assertEqual([s["page"] for s in result["segments"]], [1, 2])
        self.assertEqual(result["segments"][1]["text"], "Importo EUR 500,00")
        self.assertEqual(result["segments"][1]["method"], "ocr")

    @unittest.skipUnless(canvas is not None, "reportlab is an optional test dependency")
    def test_real_ocr_scan_when_local_tools_are_available(self):
        if not shutil.which("pdftoppm") or not shutil.which("tesseract"):
            self.skipTest("Poppler and Tesseract are optional system dependencies")
        languages = subprocess.run(
            ["tesseract", "--list-langs"], capture_output=True, text=True, timeout=15,
        )
        if "eng" not in languages.stdout.splitlines():
            self.skipTest("Tesseract English data is not installed")
        path = self.root / "real-scan.pdf"
        image = Image.new("RGB", (1500, 500), "white")
        draw = ImageDraw.Draw(image)
        font = ImageFont.load_default(size=48)
        draw.text((50, 60), "Contract amount EUR 12000", "black", font=font)
        draw.text((50, 140), "Duration twelve months", "black", font=font)
        pdf = canvas.Canvas(str(path), pagesize=(750, 250))
        pdf.drawImage(ImageReader(image), 0, 0, width=750, height=250)
        pdf.save()
        result = extract_document(path, ocr_language="eng")
        self.assertEqual(result["segments"][0]["method"], "ocr")
        self.assertIn("12000", result["segments"][0]["text"])
        self.assertIn("twelve months", result["segments"][0]["text"])

    def test_ocr_missing_binary_explains_local_dependency(self):
        with patch("documents.services.ingest.shutil.which", return_value=None):
            with self.assertRaisesRegex(IngestionError, "Poppler"):
                _ocr_page(self.root / "scan.pdf", 1, "ita+eng")

    def test_ocr_timeout_is_bounded_and_does_not_leak_subprocess_output(self):
        with patch("documents.services.ingest.shutil.which", side_effect=lambda name: f"/usr/bin/{name}"):
            with patch("documents.services.ingest.subprocess.run", side_effect=subprocess.TimeoutExpired("private-input", 90)):
                with self.assertRaisesRegex(IngestionError, "Tempo massimo OCR") as error:
                    _ocr_page(self.root / "scan.pdf", 3, "ita+eng")
        self.assertNotIn("private-input", str(error.exception))

    def test_empty_ocr_is_not_reported_as_success(self):
        path = self.root / "blank.pdf"
        writer = PdfWriter()
        writer.add_blank_page(width=200, height=200)
        writer.write(path)
        with patch("documents.services.ingest._ocr_page", return_value=""):
            with self.assertRaisesRegex(IngestionError, "Nessun testo leggibile"):
                extract_document(path)

    def test_ocr_language_does_not_allow_arbitrary_options(self):
        with self.assertRaisesRegex(IngestionError, "lingua OCR"):
            extract_document(self.make_docx(), ocr_language="ita --debug")


if __name__ == "__main__":
    unittest.main()
