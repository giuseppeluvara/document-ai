import json
from io import BytesIO
from types import SimpleNamespace
from datetime import datetime, timezone
from unittest import TestCase
from openpyxl import load_workbook
from documents.services.exports import build_xlsx

class ExportOverflowTests(TestCase):
    def test_all_long_values_evidence_and_warnings_continue_without_truncation(self):
        values = ['x' * 1000 + str(i) for i in range(40)]
        quote = 'y' * 40000
        warning = 'z' * 40000
        document = SimpleNamespace(
            id='sample', original_name='=not-a-formula.docx', sha256='a' * 64,
            created_at=datetime.now(timezone.utc), overrides={}, metrics={},
            template_snapshot={'name': 'Test', 'fields': [{'key': 'items', 'label': 'Items', 'type': 'list'}]},
            result={'engine': 'local', 'warnings': [warning], 'fields': [{'key': 'items', 'label': 'Items', 'value': values, 'status': 'found', 'evidence': [{'segment_id': 's1', 'quote': quote}], 'warnings': [warning]}]},
            extracted={'segments': [{'id': 's1', 'label': 'Blocco', 'page': None, 'method': 'docx', 'text': 'controllo\x00preservato'}], 'warnings': []})
        workbook = load_workbook(BytesIO(build_xlsx(document)))
        rows = list(workbook['Campi'].iter_rows(min_row=2, values_only=True))
        self.assertEqual(''.join(row[3] or '' for row in rows), json.dumps(values, ensure_ascii=False))
        self.assertEqual(''.join(row[6] or '' for row in rows), 's1: ' + quote)
        self.assertEqual(''.join(row[7] or '' for row in rows), warning)
        self.assertEqual(workbook['Fonti']['E2'].value, 'controllo\\u0000preservato')
        self.assertEqual(workbook['Metadati']['B1'].data_type, 's')
