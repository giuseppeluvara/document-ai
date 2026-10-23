"""Integration coverage for the local upload → queue → review → export flow.

The model response is mocked at the service boundary; DOCX ingestion, database,
request validation and both export formats use their real implementations.
"""

import hashlib
import json
import tempfile
from io import BytesIO
from pathlib import Path
from unittest.mock import patch

from docx import Document as WordDocument
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import Client, TestCase, override_settings
from django.urls import reverse
from openpyxl import load_workbook

from documents.models import Document, ExtractionTemplate
from documents.services.ai import AIError
from documents.services.exports import build_xlsx
from documents.services.jobs import claim_next, process_document


FIELDS = [
    {'key': 'oggetto', 'label': 'Oggetto', 'description': 'Oggetto del contratto.', 'type': 'text'},
    {'key': 'importo', 'label': 'Importo', 'description': 'Importo principale.', 'type': 'number'},
    {'key': 'scadenze', 'label': 'Scadenze', 'description': 'Scadenze esplicite.', 'type': 'list'},
    {'key': 'data', 'label': 'Data', 'description': 'Data della stipula.', 'type': 'date'},
]


def word_bytes():
    document = WordDocument()
    document.add_paragraph('Contratto di assistenza informatica.')
    document.add_paragraph('Oggetto: manutenzione software.')
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = 'Importo'
    table.cell(0, 1).text = '420 euro'
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def model_response(segments, fields, **kwargs):
    source = next(segment for segment in segments if 'manutenzione software' in segment['text'])
    return {
        'engine': 'test-local-model',
        'warnings': [],
        'fields': [
            {
                'key': field['key'], 'label': field['label'],
                'value': 'manutenzione software' if field['key'] == 'oggetto' else None,
                'evidence': [{'segment_id': source['id'], 'quote': 'manutenzione software'}]
                if field['key'] == 'oggetto' else [],
                'status': 'found' if field['key'] == 'oggetto' else 'not_found',
                'warnings': [],
            }
            for field in fields
        ],
    }


@override_settings(ALLOWED_HOSTS=['testserver', 'localhost', '127.0.0.1'])
class DocumentFlowTests(TestCase):
    def setUp(self):
        self.directory = tempfile.TemporaryDirectory()
        self.addCleanup(self.directory.cleanup)
        self.settings_override = override_settings(MEDIA_ROOT=Path(self.directory.name) / 'originals')
        self.settings_override.enable()
        self.addCleanup(self.settings_override.disable)
        self.template = ExtractionTemplate.objects.create(
            name='Contratto test', description='Campi configurabili', fields=FIELDS)
        self.content = word_bytes()

    def upload(self, filename='contratto.docx'):
        response = self.client.post(reverse('documents:index'), {
            'template': self.template.pk,
            'file': SimpleUploadedFile(filename, self.content, content_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document'),
        })
        self.assertEqual(response.status_code, 302)
        return Document.objects.latest('created_at')

    def complete(self):
        document = self.upload()
        claimed = claim_next()
        self.assertEqual(claimed.pk, document.pk)
        with patch('documents.services.jobs.run_extraction', side_effect=model_response):
            process_document(claimed)
        document.refresh_from_db()
        self.assertEqual(document.status, Document.Status.SUCCEEDED)
        return document

    def test_template_creation_edit_and_uploaded_snapshot(self):
        response = self.client.post(reverse('documents:template_create'), {
            'name': 'Accordo', 'description': 'Template nuovo', 'fields_json': json.dumps(FIELDS),
        })
        self.assertEqual(response.status_code, 302)
        created = ExtractionTemplate.objects.get(name='Accordo')
        self.assertEqual(created.fields, FIELDS)

        document = self.upload()
        changed = [dict(FIELDS[0], label='Oggetto aggiornato')]
        response = self.client.post(reverse('documents:template_edit', args=[self.template.pk]), {
            'name': self.template.name, 'description': 'Modificato', 'fields_json': json.dumps(changed),
        })
        self.assertEqual(response.status_code, 302)
        self.template.refresh_from_db()
        document.refresh_from_db()
        self.assertEqual(self.template.fields, changed)
        self.assertEqual(document.template_snapshot['fields'], FIELDS)

    def test_invalid_template_definitions_are_form_errors(self):
        malformed = [
            {'not': 'an array'},
            [dict(FIELDS[0], type={'unexpected': 'mapping'})],
            [dict(FIELDS[0], type=['text'])],
            [FIELDS[0], FIELDS[0]],
            [dict(FIELDS[0], key='bad-key')],
        ]
        for fields in malformed:
            with self.subTest(fields=fields):
                response = self.client.post(reverse('documents:template_create'), {
                    'name': 'Invalid', 'description': '', 'fields_json': json.dumps(fields),
                })
                self.assertEqual(response.status_code, 200)
                self.assertIn('fields_json', response.context['form'].errors)
        self.assertFalse(ExtractionTemplate.objects.filter(name='Invalid').exists())

    def test_upload_conserves_original_queues_and_uses_random_storage_name(self):
        document = self.upload('../../contratto.docx')
        self.assertEqual(document.status, Document.Status.QUEUED)
        self.assertEqual(document.sha256, hashlib.sha256(self.content).hexdigest())
        self.assertEqual(document.original_name, 'contratto.docx')
        self.assertNotIn('contratto.docx', document.original.name)
        self.assertEqual(Path(document.original.path).read_bytes(), self.content)
        status = self.client.get(reverse('documents:status', args=[document.pk]))
        self.assertEqual(status.json()['status'], 'queued')
        self.assertEqual(self.client.get(reverse('documents:export_json', args=[document.pk])).status_code, 400)

    def test_queue_ingestion_and_exports_use_real_document_sources(self):
        document = self.complete()
        self.assertIsNone(claim_next())
        self.assertGreater(document.metrics['characters'], 0)
        self.assertTrue(any('420 euro' in segment['text'] for segment in document.extracted['segments']))
        response = self.client.get(reverse('documents:detail', args=[document.pk]))
        self.assertContains(response, 'manutenzione software')
        self.assertEqual(response['Cache-Control'], 'no-store')
        payload = self.client.get(reverse('documents:export_json', args=[document.pk])).json()
        self.assertEqual(payload['document']['sha256'], document.sha256)
        self.assertEqual(payload['template']['fields'], FIELDS)
        self.assertEqual(payload['fields'][0]['value'], 'manutenzione software')
        self.assertIsNone(payload['fields'][1]['value'])
        self.assertEqual(payload['fields'][1]['status'], 'not_found')
        self.assertTrue(payload['sources'])
        response = self.client.get(reverse('documents:export_xlsx', args=[document.pk]))
        self.assertEqual(response.status_code, 200)
        workbook = load_workbook(BytesIO(response.content))
        self.assertEqual(workbook.sheetnames, ['Campi', 'Fonti', 'Metadati'])
        self.assertEqual(workbook['Campi']['D2'].value, 'manutenzione software')
        self.assertGreater(workbook['Fonti'].max_row, 1)

    def test_review_preserves_original_and_recomputes_effective_status(self):
        document = self.complete()
        response = self.client.post(reverse('documents:review', args=[document.pk]), {
            'field_oggetto': '', 'field_importo': '420',
            'field_scadenze': 'Prima scadenza\nSeconda scadenza', 'field_data': '2026-12-10',
        })
        self.assertEqual(response.status_code, 302)
        payload = self.client.get(reverse('documents:export_json', args=[document.pk])).json()
        fields = {field['key']: field for field in payload['fields']}
        self.assertIsNone(fields['oggetto']['value'])
        self.assertEqual(fields['oggetto']['status'], 'not_found')
        self.assertEqual(fields['oggetto']['automatic_status'], 'found')
        self.assertEqual(fields['oggetto']['original_value'], 'manutenzione software')
        self.assertTrue(fields['oggetto']['evidence'])
        self.assertEqual(fields['importo']['value'], 420)
        self.assertEqual(fields['importo']['status'], 'found')
        self.assertEqual(fields['importo']['automatic_status'], 'not_found')
        self.assertEqual(fields['scadenze']['value'], ['Prima scadenza', 'Seconda scadenza'])
        self.assertTrue(all(field['manually_reviewed'] for field in fields.values()))
        document.refresh_from_db()
        self.assertEqual(document.result['fields'][0]['value'], 'manutenzione software')

    def test_invalid_review_is_atomic_and_does_not_overwrite_previous_review(self):
        document = self.complete()
        valid = {'field_oggetto': 'Valore corretto', 'field_importo': '420', 'field_scadenze': '', 'field_data': ''}
        self.client.post(reverse('documents:review', args=[document.pk]), valid)
        document.refresh_from_db()
        before = document.overrides
        self.client.post(reverse('documents:review', args=[document.pk]), dict(valid, field_importo='NaN'))
        document.refresh_from_db()
        self.assertEqual(document.overrides, before)
        response = self.client.post(reverse('documents:review', args=[document.pk]), {'field_oggetto': 'Incompleto'})
        self.assertEqual(response.status_code, 400)
        document.refresh_from_db()
        self.assertEqual(document.overrides, before)

    def test_failure_retains_original_and_ingested_text_and_can_retry(self):
        document = self.upload()
        with patch('documents.services.jobs.run_extraction', side_effect=AIError('Motore locale non raggiungibile.')):
            process_document(claim_next())
        document.refresh_from_db()
        self.assertEqual(document.status, Document.Status.FAILED)
        self.assertTrue(document.extracted['segments'])
        self.assertTrue(Path(document.original.path).exists())
        self.assertEqual(document.error, 'Motore locale non raggiungibile.')
        self.client.post(reverse('documents:retry', args=[document.pk]))
        document.refresh_from_db()
        self.assertEqual(document.status, Document.Status.QUEUED)
        self.assertEqual(document.error, '')
        self.assertEqual(document.template_snapshot['fields'], FIELDS)

    def test_unexpected_failure_does_not_expose_exception_contents(self):
        document = self.upload()
        with patch('documents.services.jobs.run_extraction', side_effect=RuntimeError('PRIVATE DOCUMENT CONTENT')):
            process_document(claim_next())
        document.refresh_from_db()
        self.assertEqual(document.status, Document.Status.FAILED)
        self.assertNotIn('PRIVATE DOCUMENT CONTENT', document.error)

    def test_original_download_and_delete_remove_stored_data(self):
        document = self.complete()
        path = Path(document.original.path)
        response = self.client.get(reverse('documents:original', args=[document.pk]))
        try:
            self.assertEqual(b''.join(response.streaming_content), self.content)
            self.assertIn('attachment', response['Content-Disposition'])
        finally:
            response.close()
        with self.captureOnCommitCallbacks(execute=True):
            response = self.client.post(reverse('documents:delete', args=[document.pk]))
        self.assertEqual(response.status_code, 302)
        self.assertFalse(Document.objects.filter(pk=document.pk).exists())
        self.assertFalse(path.exists())
        self.assertTrue(ExtractionTemplate.objects.filter(pk=self.template.pk).exists())

    def test_processing_document_cannot_be_deleted_or_requeued(self):
        document = self.upload()
        claimed = claim_next()
        self.client.post(reverse('documents:delete', args=[document.pk]))
        self.client.post(reverse('documents:retry', args=[document.pk]))
        claimed.refresh_from_db()
        self.assertEqual(claimed.status, Document.Status.PROCESSING)
        self.assertTrue(Path(claimed.original.path).exists())

    def test_rejects_unsupported_and_mislabelled_files(self):
        for name, content in [('old.doc', self.content), ('fake.pdf', b'not a pdf'), ('fake.docx', b'not a zip')]:
            with self.subTest(name=name):
                response = self.client.post(reverse('documents:index'), {
                    'template': self.template.pk, 'file': SimpleUploadedFile(name, content),
                })
                self.assertEqual(response.status_code, 200)
                self.assertIn('file', response.context['form'].errors)
        self.assertEqual(Document.objects.count(), 0)

    def test_access_is_local_and_mutations_require_post_and_csrf(self):
        self.assertEqual(self.client.get(reverse('documents:index'), REMOTE_ADDR='192.0.2.9').status_code, 403)
        document = self.upload()
        for route in ['review', 'retry', 'delete']:
            self.assertEqual(self.client.get(reverse('documents:' + route, args=[document.pk])).status_code, 405)
        client = Client(enforce_csrf_checks=True)
        self.assertEqual(client.post(reverse('documents:delete', args=[document.pk])).status_code, 403)
        self.assertTrue(Document.objects.filter(pk=document.pk).exists())

    def test_excel_values_stay_literal_and_long_content_is_not_lost(self):
        document = self.complete()
        document.original_name = '=HYPERLINK("https://example.invalid")'
        long_quote = 'source-' * 7000
        document.result['fields'][0]['value'] = '=1+1'
        document.result['fields'][0]['evidence'] = [{'segment_id': 'docx-1', 'quote': long_quote}]
        workbook = load_workbook(BytesIO(build_xlsx(document)))
        self.assertEqual(workbook['Campi']['D2'].value, '=1+1')
        self.assertEqual(workbook['Campi']['D2'].data_type, 's')
        self.assertEqual(workbook['Metadati']['B1'].value, document.original_name)
        self.assertEqual(workbook['Metadati']['B1'].data_type, 's')
        all_evidence = ''.join(str(row[6].value or '') for row in workbook['Campi'].iter_rows(min_row=2))
        self.assertIn(long_quote, all_evidence)
